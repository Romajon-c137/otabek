import hashlib
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from modules.auditlog.services import record_action
from modules.catalog.models import Category, Product
from modules.documents.models import DocumentKind
from modules.documents.services import create_document, post_document
from modules.finance.models import ObligationSide
from modules.finance.services import create_obligation
from modules.identity.models import UserRole
from modules.identity.permissions import require_role, require_warehouse
from modules.parties.models import Party
from .models import ImportBatch


SHEETS = {
    "Товары": ["code", "name", "sku", "barcode", "category", "unit", "purchase_price", "retail_price", "wholesale_price"],
    "Контрагенты": ["code", "name", "kind", "phone"],
    "Остатки": ["product_code", "quantity", "unit_cost", "lot_order"],
    "Долги": ["party_code", "side", "product_code", "quantity", "unit_price", "open_amount", "reason"],
}


def make_template() -> bytes:
    book = Workbook()
    info = book.active
    info.title = "Инструкция"
    info.append(["ХозУчёт — начальная загрузка"])
    info.append(["Не меняйте названия листов и колонок. Сначала выполните проверку, затем примените пакет."])
    examples = {
        "Товары": ["P-001", "Пример товара", "SKU-001", "460000000001", "Общая", "piece", "100.00", "150.00", "130.00"],
        "Контрагенты": ["C-001", "Пример покупателя", "customer", "+996555000000"],
        "Остатки": ["P-001", "10.000", "100.00", 1],
        "Долги": ["C-001", "receivable", "P-001", "2.000", "150.00", "300.00", "Начальный долг"],
    }
    for name, headers in SHEETS.items():
        sheet = book.create_sheet(name)
        sheet.append(headers)
        sheet.append(examples[name])
        sheet.freeze_panes = "A2"
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def _text(value):
    return str(value or "").strip()


def _number(value, *, sheet, row, field, errors, positive=False, allow_blank=False):
    raw = _text(value)
    if allow_blank and not raw:
        return ""
    try:
        number = Decimal(raw)
        if number < 0 or (positive and number <= 0):
            raise InvalidOperation
        return format(number, "f")
    except (InvalidOperation, ValueError):
        errors.append({"sheet": sheet, "row": row, "field": field, "message": "Ожидается положительное число." if positive else "Ожидается неотрицательное число."})
        return "0"


def _sheet_rows(workbook, name, errors):
    if name not in workbook.sheetnames:
        errors.append({"sheet": name, "row": 1, "message": "Обязательный лист отсутствует."})
        return []
    iterator = workbook[name].iter_rows(values_only=True)
    if [_text(value) for value in next(iterator, [])] != SHEETS[name]:
        errors.append({"sheet": name, "row": 1, "message": "Заголовки не соответствуют шаблону."})
        return []
    return [(index, dict(zip(SHEETS[name], values))) for index, values in enumerate(iterator, start=2) if any(_text(value) for value in values)]


def validate_workbook(file, actor) -> ImportBatch:
    require_role(actor, UserRole.OWNER, UserRole.MANAGER)
    raw = file.read()
    if len(raw) > 5 * 1024 * 1024:
        raise ValueError("Файл превышает 5 МБ.")
    digest = hashlib.sha256(raw).hexdigest()
    try:
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Файл не является корректной книгой XLSX.") from exc
    errors = []
    parsed = {key: [] for key in ("products", "parties", "balances", "debts")}

    product_codes = set()
    product_units = {}
    for row_number, source in _sheet_rows(workbook, "Товары", errors):
        row = {key: _text(value) for key, value in source.items()}
        if not row["code"] or not row["name"]:
            errors.append({"sheet": "Товары", "row": row_number, "message": "Код и название обязательны."})
        if row["code"] in product_codes:
            errors.append({"sheet": "Товары", "row": row_number, "field": "code", "message": "Код повторяется в файле."})
        product_codes.add(row["code"])
        product_units[row["code"]] = row["unit"]
        if row["unit"] not in Product.Unit.values:
            errors.append({"sheet": "Товары", "row": row_number, "field": "unit", "message": "Допустимо: piece, kg или meter."})
        current = Product.objects.filter(sku=row["code"]).first()
        if current and current.unit != row["unit"] and current.stock_movements.exists():
            errors.append({"sheet": "Товары", "row": row_number, "field": "unit", "message": "Единицу нельзя менять после складских движений."})
        for field in ("purchase_price", "retail_price", "wholesale_price"):
            row[field] = _number(source[field], sheet="Товары", row=row_number, field=field, errors=errors)
        row["row"] = row_number
        parsed["products"].append(row)

    party_codes = set()
    for row_number, source in _sheet_rows(workbook, "Контрагенты", errors):
        row = {key: _text(value) for key, value in source.items()}
        if not row["code"] or not row["name"]:
            errors.append({"sheet": "Контрагенты", "row": row_number, "message": "Код и название обязательны."})
        if row["code"] in party_codes:
            errors.append({"sheet": "Контрагенты", "row": row_number, "field": "code", "message": "Код повторяется в файле."})
        party_codes.add(row["code"])
        if row["kind"] not in Party.Kind.values:
            errors.append({"sheet": "Контрагенты", "row": row_number, "field": "kind", "message": "Допустимо: customer, supplier или both."})
        row["row"] = row_number
        parsed["parties"].append(row)

    for row_number, source in _sheet_rows(workbook, "Остатки", errors):
        row = {key: _text(value) for key, value in source.items()}
        if row["product_code"] not in product_codes and not Product.objects.filter(sku=row["product_code"]).exists():
            errors.append({"sheet": "Остатки", "row": row_number, "field": "product_code", "message": "Товар с таким кодом не найден."})
        row["quantity"] = _number(source["quantity"], sheet="Остатки", row=row_number, field="quantity", errors=errors, positive=True)
        unit = product_units.get(row["product_code"]) or Product.objects.filter(sku=row["product_code"]).values_list("unit", flat=True).first()
        if unit == Product.Unit.PIECE and Decimal(row["quantity"]) != Decimal(row["quantity"]).to_integral_value():
            errors.append({"sheet": "Остатки", "row": row_number, "field": "quantity", "message": "Для штучного товара количество должно быть целым."})
        row["unit_cost"] = _number(source["unit_cost"], sheet="Остатки", row=row_number, field="unit_cost", errors=errors)
        row["lot_order"] = _text(source["lot_order"] or row_number)
        row["row"] = row_number
        parsed["balances"].append(row)

    for row_number, source in _sheet_rows(workbook, "Долги", errors):
        row = {key: _text(value) for key, value in source.items()}
        if row["party_code"] not in party_codes and not Party.objects.filter(code=row["party_code"]).exists():
            errors.append({"sheet": "Долги", "row": row_number, "field": "party_code", "message": "Контрагент с таким кодом не найден."})
        if row["side"] not in ObligationSide.values:
            errors.append({"sheet": "Долги", "row": row_number, "field": "side", "message": "Допустимо: receivable или payable."})
        if row["product_code"] and row["product_code"] not in product_codes and not Product.objects.filter(sku=row["product_code"]).exists():
            errors.append({"sheet": "Долги", "row": row_number, "field": "product_code", "message": "Товар с таким кодом не найден."})
        row["quantity"] = _number(source["quantity"], sheet="Долги", row=row_number, field="quantity", errors=errors, positive=True, allow_blank=True)
        unit = product_units.get(row["product_code"]) or (Product.objects.filter(sku=row["product_code"]).values_list("unit", flat=True).first() if row["product_code"] else None)
        if row["quantity"] and unit == Product.Unit.PIECE and Decimal(row["quantity"]) != Decimal(row["quantity"]).to_integral_value():
            errors.append({"sheet": "Долги", "row": row_number, "field": "quantity", "message": "Для штучного товара количество должно быть целым."})
        row["unit_price"] = _number(source["unit_price"], sheet="Долги", row=row_number, field="unit_price", errors=errors, allow_blank=True)
        row["open_amount"] = _number(source["open_amount"], sheet="Долги", row=row_number, field="open_amount", errors=errors, positive=True)
        detail_flags = [bool(row["product_code"]), bool(row["quantity"]), bool(row["unit_price"])]
        if any(detail_flags) and not all(detail_flags):
            errors.append({"sheet": "Долги", "row": row_number, "message": "Для товарного долга одновременно заполните товар, количество и цену."})
        row["row"] = row_number
        parsed["debts"].append(row)
    return ImportBatch.objects.create(source_name=file.name[:240], content_hash=digest, rows=parsed, errors=errors, status=ImportBatch.Status.VALIDATED if not errors else ImportBatch.Status.REJECTED, created_by=actor)


@transaction.atomic
def apply_batch(batch: ImportBatch, actor, warehouse) -> dict:
    require_role(actor, UserRole.OWNER, UserRole.MANAGER)
    require_warehouse(actor, warehouse, "can_manage_stock")
    batch = ImportBatch.objects.select_for_update().get(pk=batch.pk)
    if batch.created_by_id != actor.id and not actor.is_superuser:
        from rest_framework.exceptions import PermissionDenied
        raise PermissionDenied("Нельзя применить чужой импорт.")
    if batch.status != ImportBatch.Status.VALIDATED:
        raise ValueError("Этот импорт нельзя применить.")
    if ImportBatch.objects.filter(content_hash=batch.content_hash, status=ImportBatch.Status.APPLIED).exclude(pk=batch.pk).exists():
        raise ValueError("Этот файл уже был применён.")

    products = {}
    product_created = product_updated = 0
    for row in batch.rows["products"]:
        category = None
        if row["category"]:
            category, _ = Category.objects.get_or_create(name=row["category"])
        product, was_created = Product.objects.update_or_create(
            sku=row["code"],
            defaults={"name": row["name"], "barcode": row["barcode"], "category": category, "unit": row["unit"], "purchase_price": row["purchase_price"], "retail_price": row["retail_price"], "wholesale_price": row["wholesale_price"], "is_active": True},
        )
        products[row["code"]] = product
        product_created += int(was_created)
        product_updated += int(not was_created)

    parties = {}
    party_created = party_updated = 0
    for row in batch.rows["parties"]:
        party, was_created = Party.objects.update_or_create(code=row["code"], defaults={"name": row["name"], "kind": row["kind"], "phone": row["phone"], "is_active": True})
        parties[row["code"]] = party
        party_created += int(was_created)
        party_updated += int(not was_created)

    balance_lines = []
    for row in sorted(batch.rows["balances"], key=lambda item: (item["lot_order"], item["row"])):
        product = products.get(row["product_code"]) or Product.objects.get(sku=row["product_code"])
        balance_lines.append({"product": product, "quantity": row["quantity"], "unit_price": row["unit_cost"], "discount_amount": "0.00"})
    if balance_lines:
        opening = create_document(actor=actor, warehouse=warehouse, kind=DocumentKind.OPENING, effective_date=timezone.localdate(), note=f"Начальные остатки: {batch.source_name}", lines=balance_lines)
        post_document(actor=actor, document=opening, idempotency_key=f"import-{batch.content_hash}-opening")

    for row in batch.rows["debts"]:
        party = parties.get(row["party_code"]) or Party.objects.get(code=row["party_code"])
        product = products.get(row["product_code"]) or (Product.objects.get(sku=row["product_code"]) if row["product_code"] else None)
        create_obligation(
            party=party, warehouse=warehouse, side=row["side"], amount=Decimal(row["open_amount"]), effective_date=timezone.localdate(), actor=actor,
            product=product, original_quantity=Decimal(row["quantity"]) if row["quantity"] else None,
            unit_price=Decimal(row["unit_price"]) if row["unit_price"] else None,
            reason=row["reason"] or f"Начальный долг: {batch.source_name}",
        )

    batch.status = ImportBatch.Status.APPLIED
    batch.applied_at = timezone.now()
    batch.save(update_fields=["status", "applied_at"])
    result = {"products_created": product_created, "products_updated": product_updated, "parties_created": party_created, "parties_updated": party_updated, "stock_lines": len(balance_lines), "debts": len(batch.rows["debts"])}
    record_action(actor=actor, action="import.applied", entity=batch, warehouse_id=warehouse.pk, after=result)
    return result

