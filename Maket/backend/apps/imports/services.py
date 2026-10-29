import hashlib
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from apps.audit.models import AuditEvent
from apps.catalog.models import Category, Product
from apps.documents.models import Document, OpeningStock, OpeningStockLine
from apps.documents.services import post_opening_stock
from apps.partners.models import Counterparty
from apps.settlements.models import Obligation, SettlementMovement
from .models import ImportBatch


UNIT_MAP = {"шт": "pcs", "pcs": "pcs", "кг": "kg", "kg": "kg", "м": "m", "m": "m"}
SIDE_MAP = {
    "нам должны": Obligation.Side.RECEIVABLE, "receivable": Obligation.Side.RECEIVABLE,
    "мы должны": Obligation.Side.PAYABLE, "payable": Obligation.Side.PAYABLE,
}


def _decimal(value, *, sheet, row, field, errors, positive=False, nonnegative=False):
    try:
        result = Decimal(str(value).replace(",", "."))
    except (InvalidOperation, TypeError, ValueError):
        errors.append({"sheet": sheet, "row": row, "field": field, "message": "Ожидается число."})
        return None
    if positive and result <= 0:
        errors.append({"sheet": sheet, "row": row, "field": field, "message": "Значение должно быть больше нуля."})
    if nonnegative and result < 0:
        errors.append({"sheet": sheet, "row": row, "field": field, "message": "Значение не может быть отрицательным."})
    return result


def _rows(workbook, sheet_name, required, errors):
    if sheet_name not in workbook.sheetnames:
        errors.append({"sheet": sheet_name, "row": 1, "field": "sheet", "message": "Отсутствует обязательный лист."})
        return []
    sheet = workbook[sheet_name]
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        errors.append({"sheet": sheet_name, "row": 1, "field": "header", "message": "Лист пуст."})
        return []
    headers = [str(value).strip() if value is not None else "" for value in values[0]]
    missing = [name for name in required if name not in headers]
    for name in missing:
        errors.append({"sheet": sheet_name, "row": 1, "field": name, "message": "Нет обязательной колонки."})
    if missing:
        return []
    result = []
    for row_number, row in enumerate(values[1:], start=2):
        if all(value is None or str(value).strip() == "" for value in row):
            continue
        result.append((row_number, {header: row[index] if index < len(row) else None for index, header in enumerate(headers)}))
    return result


def validate_workbook(*, content: bytes, file_name: str, warehouse, actor):
    checksum = hashlib.sha256(content).hexdigest()
    errors = []
    payload = {"products": [], "counterparties": [], "stocks": [], "debts": []}
    if len(content) > 10 * 1024 * 1024:
        errors.append({"sheet": "file", "row": 0, "field": "size", "message": "Файл больше 10 МБ."})
    else:
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except Exception:
            errors.append({"sheet": "file", "row": 0, "field": "format", "message": "Не удалось прочитать файл Excel."})
            workbook = None
        if workbook:
            product_codes = set()
            for row_number, row in _rows(workbook, "Товары", ["Код", "Наименование", "Единица"], errors):
                code = str(row["Код"] or "").strip()
                name = str(row["Наименование"] or "").strip()
                unit_raw = str(row["Единица"] or "").strip().lower()
                unit = UNIT_MAP.get(unit_raw)
                if not code:
                    errors.append({"sheet": "Товары", "row": row_number, "field": "Код", "message": "Код обязателен."})
                if code in product_codes:
                    errors.append({"sheet": "Товары", "row": row_number, "field": "Код", "message": "Дублирующийся код."})
                if not name:
                    errors.append({"sheet": "Товары", "row": row_number, "field": "Наименование", "message": "Наименование обязательно."})
                if not unit:
                    errors.append({"sheet": "Товары", "row": row_number, "field": "Единица", "message": "Допустимы шт, кг или м."})
                prices = {}
                for column, key in [("Закупочная цена", "purchase_price"), ("Розничная цена", "retail_price"), ("Оптовая цена", "wholesale_price")]:
                    value = row.get(column) if row.get(column) not in (None, "") else 0
                    parsed = _decimal(value, sheet="Товары", row=row_number, field=column, errors=errors, nonnegative=True)
                    prices[key] = str(parsed) if parsed is not None else "0"
                product_codes.add(code)
                payload["products"].append({
                    "code": code, "name": name, "unit": unit or unit_raw,
                    "category": str(row.get("Категория") or "").strip(),
                    **prices,
                })

            cp_codes = set()
            for row_number, row in _rows(workbook, "Контрагенты", ["Код", "Наименование"], errors):
                code = str(row["Код"] or "").strip(); name = str(row["Наименование"] or "").strip()
                if not code or code in cp_codes:
                    errors.append({"sheet": "Контрагенты", "row": row_number, "field": "Код", "message": "Код обязателен и не должен повторяться."})
                if not name:
                    errors.append({"sheet": "Контрагенты", "row": row_number, "field": "Наименование", "message": "Наименование обязательно."})
                cp_codes.add(code)
                payload["counterparties"].append({"code": code, "name": name, "phone": str(row.get("Телефон") or "").strip()})

            for row_number, row in _rows(workbook, "Остатки", ["Код товара", "Количество", "Стоимость единицы"], errors):
                code = str(row["Код товара"] or "").strip()
                qty = _decimal(row["Количество"], sheet="Остатки", row=row_number, field="Количество", errors=errors, positive=True)
                cost = _decimal(row["Стоимость единицы"], sheet="Остатки", row=row_number, field="Стоимость единицы", errors=errors, nonnegative=True)
                if code not in product_codes and not Product.objects.filter(code=code).exists():
                    errors.append({"sheet": "Остатки", "row": row_number, "field": "Код товара", "message": "Товар не найден."})
                payload["stocks"].append({"product_code": code, "quantity": str(qty or 0), "unit_cost": str(cost or 0)})

            for row_number, row in _rows(workbook, "Долги", ["Код контрагента", "Направление", "Сумма"], errors):
                code = str(row["Код контрагента"] or "").strip()
                side_raw = str(row["Направление"] or "").strip().lower(); side = SIDE_MAP.get(side_raw)
                amount = _decimal(row["Сумма"], sheet="Долги", row=row_number, field="Сумма", errors=errors, positive=True)
                if code not in cp_codes and not Counterparty.objects.filter(code=code).exists():
                    errors.append({"sheet": "Долги", "row": row_number, "field": "Код контрагента", "message": "Контрагент не найден."})
                if not side:
                    errors.append({"sheet": "Долги", "row": row_number, "field": "Направление", "message": "Укажите «нам должны» или «мы должны»."})
                payload["debts"].append({"counterparty_code": code, "side": side or side_raw, "amount": str(amount or 0)})

            existing_products = set(Product.objects.filter(code__in=product_codes).values_list("code", flat=True))
            existing_cps = set(Counterparty.objects.filter(code__in=cp_codes).values_list("code", flat=True))
            for code in sorted(existing_products):
                errors.append({"sheet": "Товары", "row": 0, "field": "Код", "message": f"Товар {code} уже существует."})
            for code in sorted(existing_cps):
                errors.append({"sheet": "Контрагенты", "row": 0, "field": "Код", "message": f"Контрагент {code} уже существует."})

    return ImportBatch.objects.create(
        checksum=checksum, file_name=file_name[:255],
        state=ImportBatch.State.INVALID if errors else ImportBatch.State.VALID,
        warehouse=warehouse, payload=payload, errors=errors, created_by=actor,
    )


@transaction.atomic
def apply_batch(*, batch_id, checksum, actor):
    batch = ImportBatch.objects.select_for_update().get(pk=batch_id)
    if batch.created_by_id != actor.id and not actor.is_superuser:
        raise PermissionError("Чужой пакет импорта недоступен.")
    if checksum != batch.checksum:
        raise ValueError("Контрольная сумма не совпадает с проверенным файлом.")
    if batch.state != ImportBatch.State.VALID:
        raise ValueError("Можно применить только успешно проверенный пакет.")
    if ImportBatch.objects.select_for_update().filter(checksum=batch.checksum, state=ImportBatch.State.APPLIED).exclude(pk=batch.pk).exists():
        raise ValueError("Этот файл уже применён.")

    categories = {}
    for row in batch.payload["products"]:
        name = " ".join((row.get("category") or "").split())
        if name and name not in categories:
            categories[name], _ = Category.objects.get_or_create(name=name)
    Product.objects.bulk_create([
        Product(**{key: value for key, value in row.items() if key != "category"},
                category=categories.get(" ".join((row.get("category") or "").split())))
        for row in batch.payload["products"]
    ])
    Counterparty.objects.bulk_create([Counterparty(**row) for row in batch.payload["counterparties"]])
    products = {item.code: item for item in Product.objects.filter(code__in=[row["product_code"] for row in batch.payload["stocks"]])}
    counterparties = {item.code: item for item in Counterparty.objects.filter(code__in=[row["counterparty_code"] for row in batch.payload["debts"]])}

    related = []
    if batch.payload["stocks"]:
        document = Document.objects.create(
            type=Document.Type.OPENING_STOCK, number=f"IMP-{batch.id.hex[:12].upper()}",
            warehouse=batch.warehouse, author=actor, reason=f"Импорт {batch.file_name}",
        )
        opening = OpeningStock.objects.create(document=document)
        for position, row in enumerate(batch.payload["stocks"], start=1):
            product = products[row["product_code"]]
            qty = Decimal(row["quantity"]); price = Decimal(row["unit_cost"])
            OpeningStockLine.objects.create(
                opening_stock=opening, position=position, product=product,
                product_code=product.code, product_name=product.name, unit=product.unit,
                quantity=qty, unit_cost=price, total_cost=qty * price,
            )
        post_opening_stock(
            opening_id=opening.pk, actor=actor, version=1,
            idempotency_key=f"import:{batch.checksum}:stock",
        )
        related.append(str(document.id))

    for position, row in enumerate(batch.payload["debts"], start=1):
        cp = counterparties[row["counterparty_code"]]
        amount = Decimal(row["amount"])
        document = Document.objects.create(
            type=Document.Type.OPENING_DEBT,
            number=f"IMP-DEBT-{batch.id.hex[:8].upper()}-{position}",
            status=Document.Status.POSTED, version=2,
            counterparty=cp, author=actor, posted_at=timezone.now(), reason=f"Импорт {batch.file_name}",
        )
        obligation = Obligation.objects.create(
            counterparty=cp, source_document=document, side=row["side"],
            original_amount=amount, open_amount=amount,
        )
        signed = amount if row["side"] == Obligation.Side.RECEIVABLE else -amount
        SettlementMovement.objects.create(
            document=document, counterparty=cp, obligation=obligation, signed_amount=signed,
        )
        related.append(str(document.id))

    batch.state = ImportBatch.State.APPLIED
    batch.applied_at = timezone.now()
    batch.save(update_fields=["state", "applied_at"])
    AuditEvent.objects.create(
        actor=actor, action="import.applied",
        payload={"batch_id": str(batch.id), "checksum": batch.checksum, "documents": related},
    )
    return {"batch_id": str(batch.id), "state": batch.state, "related_documents": related}
