from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook
from rest_framework import status
from rest_framework.test import APITestCase

from apps.catalog.models import Product
from apps.inventory.models import StockBalance
from apps.partners.models import Counterparty
from apps.settlements.models import Obligation
from apps.warehouses.models import Warehouse, WarehouseAccess


def workbook_bytes(unit="шт"):
    workbook = Workbook()
    products = workbook.active; products.title = "Товары"
    products.append(["Код", "Наименование", "Единица", "Закупочная цена", "Розничная цена", "Оптовая цена"])
    products.append(["IMP-P", "Импортный товар", unit, 100, 150, 140])
    cps = workbook.create_sheet("Контрагенты")
    cps.append(["Код", "Наименование", "Телефон"]); cps.append(["IMP-C", "Импортный клиент", "123"])
    stocks = workbook.create_sheet("Остатки")
    stocks.append(["Код товара", "Количество", "Стоимость единицы"]); stocks.append(["IMP-P", 5, 100])
    debts = workbook.create_sheet("Долги")
    debts.append(["Код контрагента", "Направление", "Сумма"]); debts.append(["IMP-C", "нам должны", 250])
    output = BytesIO(); workbook.save(output)
    return output.getvalue()


class ImportWorkflowTests(APITestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="import-admin", password="test")
        self.warehouse = Warehouse.objects.create(code="IMPORT", name="Склад импорта")
        WarehouseAccess.objects.create(user=self.user, warehouse=self.warehouse, role=WarehouseAccess.Role.ADMIN)
        self.client.force_authenticate(self.user)

    def validate(self, content):
        upload = SimpleUploadedFile(
            "initial.xlsx", content,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        return self.client.post("/api/v1/imports/validate/", {
            "warehouse": str(self.warehouse.id), "file": upload,
        }, format="multipart")

    def test_invalid_unit_reports_sheet_row_field_and_writes_nothing(self):
        response = self.validate(workbook_bytes(unit="коробка"))
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["state"], "invalid")
        error = next(item for item in response.data["errors"] if item["field"] == "Единица")
        self.assertEqual((error["sheet"], error["row"]), ("Товары", 2))
        self.assertFalse(Product.objects.filter(code="IMP-P").exists())

    def test_valid_file_applies_atomically_and_cannot_be_applied_twice(self):
        checked = self.validate(workbook_bytes())
        self.assertEqual(checked.data["state"], "valid", checked.data["errors"])
        applied = self.client.post(
            f"/api/v1/imports/{checked.data['batch_id']}/apply/",
            {"checksum": checked.data["checksum"]}, format="json",
        )
        self.assertEqual(applied.status_code, status.HTTP_200_OK, applied.data)
        product = Product.objects.get(code="IMP-P")
        self.assertEqual(StockBalance.objects.get(product=product).quantity, 5)
        cp = Counterparty.objects.get(code="IMP-C")
        self.assertEqual(Obligation.objects.get(counterparty=cp).open_amount, 250)

        repeated = self.client.post(
            f"/api/v1/imports/{checked.data['batch_id']}/apply/",
            {"checksum": checked.data["checksum"]}, format="json",
        )
        self.assertEqual(repeated.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(Product.objects.filter(code="IMP-P").count(), 1)
