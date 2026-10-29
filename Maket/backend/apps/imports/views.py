from io import BytesIO

from django.http import HttpResponse
from drf_spectacular.utils import OpenApiTypes, extend_schema
from openpyxl import Workbook
from rest_framework import status
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.warehouses.models import WarehouseAccess
from .serializers import ApplyImportSerializer, ValidateImportSerializer
from .services import apply_batch, validate_workbook


def ensure_admin(user, warehouse):
    if user.is_superuser:
        return
    if not WarehouseAccess.objects.filter(
        user=user, warehouse=warehouse, role=WarehouseAccess.Role.ADMIN,
    ).exists():
        raise PermissionDenied("Импорт доступен администратору склада.")


class ValidateImportView(APIView):
    @extend_schema(request=ValidateImportSerializer, responses=OpenApiTypes.OBJECT)
    def post(self, request):
        serializer = ValidateImportSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        warehouse = serializer.validated_data["warehouse"]
        ensure_admin(request.user, warehouse)
        uploaded = serializer.validated_data["file"]
        batch = validate_workbook(
            content=uploaded.read(), file_name=uploaded.name,
            warehouse=warehouse, actor=request.user,
        )
        return Response({
            "batch_id": str(batch.id), "checksum": batch.checksum,
            "state": batch.state, "counts": {key: len(value) for key, value in batch.payload.items()},
            "errors": batch.errors,
        }, status=status.HTTP_200_OK)


class ApplyImportView(APIView):
    @extend_schema(request=ApplyImportSerializer, responses=OpenApiTypes.OBJECT)
    def post(self, request, batch_id):
        serializer = ApplyImportSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        from .models import ImportBatch
        try:
            batch = ImportBatch.objects.get(pk=batch_id)
        except ImportBatch.DoesNotExist:
            return Response({"detail": "Пакет не найден."}, status=404)
        ensure_admin(request.user, batch.warehouse)
        try:
            result = apply_batch(
                batch_id=batch_id, checksum=serializer.validated_data["checksum"], actor=request.user,
            )
        except (ValueError, PermissionError) as exc:
            raise ValidationError({"detail": str(exc)}) from exc
        return Response(result)


class ImportTemplateView(APIView):
    @extend_schema(responses={(200, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"): bytes})
    def get(self, request):
        workbook = Workbook()
        products = workbook.active; products.title = "Товары"
        products.append(["Код", "Наименование", "Единица", "Закупочная цена", "Розничная цена", "Оптовая цена", "Артикул", "Штрихкод"])
        counterparties = workbook.create_sheet("Контрагенты")
        counterparties.append(["Код", "Наименование", "Телефон"])
        stocks = workbook.create_sheet("Остатки")
        stocks.append(["Код товара", "Количество", "Стоимость единицы"])
        debts = workbook.create_sheet("Долги")
        debts.append(["Код контрагента", "Направление", "Сумма"])
        output = BytesIO(); workbook.save(output)
        response = HttpResponse(
            output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="hozuchet-import-template.xlsx"'
        return response
